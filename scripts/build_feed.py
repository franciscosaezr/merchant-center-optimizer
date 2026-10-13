#!/usr/bin/env python3
"""
build_feed.py — Construye un feed suplementario de Merchant Center para Google Sheets.

Uso:
    python build_feed.py productos.json salida/feed_suplementario.xlsx
    python build_feed.py productos.json salida/feed_suplementario.xlsx --sheets-json salida/sheets.json

Entrada: JSON con valores LIMPIOS (sin escapar). Estructura en references/formato-sheets.md.
Salida: .xlsx con dos pestañas:
    - Feed:  la que lee Merchant Center. Columnas repetidas, escapado aplicado, celdas en texto.
    - Notas: avisos automáticos + notas del JSON.
Con --sheets-json, además un JSON con los valores y las peticiones batchUpdate para
crear el feed directamente en Google Sheets con el conector (ver references/google-sheets.md).

Requiere: openpyxl
"""
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: pip install openpyxl")

AI_SOURCE = "trained_algorithmic_media"

# Orden de columnas en la pestaña Feed. (nombre, tipo)
#   simple  -> una columna, sin escapado
#   ai      -> structured_title / structured_description
#   list    -> repetido simple (una columna por valor, se escapan comas)
#   url     -> repetido de URLs (comas codificadas como %2C)
#   detail / qa / related / variant -> repetidos de grupo (se escapan comas y dos puntos)
COLUMNS = [
    ("id", "simple"),
    ("feed_label", "simple"),
    ("content_language", "simple"),
    ("title", "simple"),
    ("structured_title", "ai"),
    ("short_title", "simple"),
    ("description", "simple"),
    ("structured_description", "ai"),
    ("google_product_category", "simple"),
    ("product_type", "simple"),
    ("brand", "simple"),
    ("gtin", "simple"),
    ("mpn", "simple"),
    ("item_group_title", "simple"),
    ("variant_option", "variant"),
    ("product_highlight", "list"),
    ("product_detail", "detail"),
    ("question_and_answer", "qa"),
    ("related_product", "related"),
    ("popularity_rank", "simple"),
    ("document_link", "url"),
    ("video_link", "url"),
]

CONVERSATIONAL = {
    "question_and_answer", "related_product", "popularity_rank", "document_link",
    "item_group_title", "variant_option", "video_link",
}

PRIMARY_USUALLY_SENDS = {
    "title", "description", "brand", "gtin", "mpn",
    "google_product_category", "product_type",
}

REL_TYPES = {"part_of_set", "required_part", "often_bought_with",
             "substitute", "different_brand", "accessory"}

REPEATED = {"list", "url", "detail", "qa", "related", "variant"}


class BuildError(Exception):
    pass


def esc_list(text: str) -> str:
    """Escapado para atributos repetidos simples en Google Sheets."""
    return str(text).replace("\\", "\\\\").replace(",", "\\,")


def esc_group(text: str) -> str:
    """Escapado para subatributos de atributos de grupo en Google Sheets."""
    return str(text).replace("\\", "\\\\").replace(",", "\\,").replace(":", "\\:")


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def fmt_ai(content: str, pid: str, attr: str) -> str:
    content = clean(content)
    if not content:
        return ""
    if '"' in content:
        raise BuildError(
            f"[{pid}] {attr}: el texto contiene comillas dobles y rompe el formato "
            f"trained_algorithmic_media:\"...\". Quítalas o sustitúyelas."
        )
    return f'{AI_SOURCE}:"{content}"'


def fmt_rank(value, pid: str) -> str:
    if value in (None, ""):
        return ""
    try:
        num = float(str(value).replace(",", "."))
    except ValueError:
        raise BuildError(f"[{pid}] popularity_rank no es un número: {value!r}")
    if not 0 <= num <= 100:
        raise BuildError(f"[{pid}] popularity_rank fuera de rango 0–100: {value!r}")
    num = round(num, 1)
    return str(int(num)) if num == int(num) else f"{num:.1f}"


def items_for(product: dict, attr: str, kind: str) -> list:
    """Devuelve la lista de valores ya formateados y escapados para un atributo repetido."""
    raw = product.get(attr) or []
    if isinstance(raw, (str, dict)):
        raw = [raw]
    pid = product.get("id", "?")
    out = []
    for item in raw:
        if kind == "list":
            v = clean(item)
            if v:
                out.append(esc_list(v))
        elif kind == "url":
            v = str(item).strip().replace(",", "%2C").replace(" ", "%20")
            if v:
                out.append(v)
        elif kind == "detail":
            sec = clean(item.get("section", ""))
            name = clean(item.get("name", ""))
            val = clean(item.get("value", ""))
            if not name or not val:
                raise BuildError(f"[{pid}] product_detail sin name o value: {item}")
            out.append(f"{esc_group(sec)}:{esc_group(name)}:{esc_group(val)}")
        elif kind == "qa":
            q = clean(item.get("question", ""))
            a = clean(item.get("answer", ""))
            if not q or not a:
                raise BuildError(f"[{pid}] question_and_answer sin pregunta o respuesta: {item}")
            out.append(f"{esc_group(q)}:{esc_group(a)}")
        elif kind == "related":
            rt = clean(item.get("relationship_type", ""))
            it = clean(item.get("identifier_type", ""))
            ident = clean(item.get("identifier", ""))
            if not (rt and it and ident):
                raise BuildError(f"[{pid}] related_product incompleto: {item}")
            if rt not in REL_TYPES:
                raise BuildError(f"[{pid}] related_product: relationship_type '{rt}' no válido. "
                                 f"Valores: {', '.join(sorted(REL_TYPES))}")
            if it not in ("gtin", "id"):
                raise BuildError(f"[{pid}] related_product: identifier_type '{it}' no válido (gtin o id).")
            if not re.fullmatch(r"[A-Za-z0-9_-]+", ident):
                raise BuildError(f"[{pid}] related_product: identificador '{ident}' con caracteres "
                                 f"no permitidos (solo letras, números, _ y -).")
            out.append(f"{rt}:{it}:{ident}")
        elif kind == "variant":
            n = clean(item.get("name", ""))
            v = clean(item.get("value", ""))
            if not (n and v):
                raise BuildError(f"[{pid}] variant_option incompleto: {item}")
            out.append(f"{esc_group(n)}:{esc_group(v)}")
    return out


def single_value(product: dict, attr: str, kind: str) -> str:
    pid = product.get("id", "?")
    if kind == "ai":
        return fmt_ai(product.get(attr), pid, attr)
    if attr == "popularity_rank":
        return fmt_rank(product.get(attr), pid)
    return clean(product.get(attr))


def build(data: dict):
    products = data.get("productos") or data.get("products") or []
    if not products:
        raise BuildError("El JSON no tiene productos (clave 'productos').")

    ids = [clean(p.get("id")) for p in products]
    if any(not i for i in ids):
        raise BuildError("Hay productos sin id.")
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        raise BuildError(f"IDs duplicados: {', '.join(sorted(dup))}")

    # Calcular valores y número de columnas por atributo
    values = []  # por producto: {attr: str | list}
    width = {}
    for p in products:
        row = {}
        for attr, kind in COLUMNS:
            if kind in REPEATED:
                row[attr] = items_for(p, attr, kind)
                width[attr] = max(width.get(attr, 0), len(row[attr]))
            else:
                row[attr] = single_value(p, attr, kind)
                if row[attr]:
                    width[attr] = max(width.get(attr, 0), 1)
        values.append(row)
    width["id"] = 1

    headers, header_attrs = [], []
    for attr, kind in COLUMNS:
        n = width.get(attr, 0)
        headers.extend([attr] * n)
        header_attrs.extend([(attr, kind, k) for k in range(n)])

    # Notas automáticas
    notes = []
    uses_plain = any(values[i]["title"] or values[i]["description"] for i in range(len(products)))
    if uses_plain:
        notes.append(
            "AVISO CONTENIDO IA: los títulos y descripciones de este feed se han redactado con ayuda de IA "
            "y se envían en title y description. Google pide declarar el contenido generado con IA en "
            "structured_title y structured_description; enviarlo en title y description puede generar "
            "advertencias o desaprobaciones. Revísalos y edítalos antes de subir, y vigila la sección de "
            "problemas de Merchant Center tras la carga."
        )
    uses_ai = any(values[i]["structured_title"] or values[i]["structured_description"]
                  for i in range(len(products)))
    if uses_ai:
        notes.append(
            "ANTES DE DAR POR BUENO EL FEED: si el feed principal envía title o description, "
            "Google usa esos y no los structured_title / structured_description de este archivo. "
            "Crea en Merchant Center una regla de atributo que vacíe (Borrar / Clear) title y "
            "description para estos productos, o deja de enviarlos desde la herramienta de feeds. "
            "Después comprueba en el diagnóstico de un producto qué título se está usando."
        )
    for i, p in enumerate(products):
        r = values[i]
        if r["title"] and r["structured_title"]:
            notes.append(f"[{ids[i]}] Lleva title y structured_title: Google usará title.")
        if r["description"] and r["structured_description"]:
            notes.append(f"[{ids[i]}] Lleva description y structured_description: Google usará description.")

    for attr, kind in COLUMNS:
        if attr not in PRIMARY_USUALLY_SENDS or not width.get(attr):
            continue
        empty = [ids[i] for i, r in enumerate(values) if not r[attr]]
        if empty:
            notes.append(
                f"Columna {attr} con celdas vacías en: {', '.join(empty)}. No está documentado si una "
                f"celda vacía en el suplementario borra el valor del principal. Rellénala o saca esos "
                f"productos a otro archivo."
            )

    notes.append(
        "Pestaña 'Feed' = la que tiene que leer Merchant Center como fuente suplementaria. "
        "Cabeceras en verde = atributos conversacionales."
    )
    notes.append(
        "Escapado (Google Sheets): en los atributos repetidos y de grupo, las comas del texto van "
        "como \\, y los dos puntos del texto como \\: . Los dos puntos sin escapar separan "
        "subatributos. Si tras subirlo ves valores partidos, revisa esto primero."
    )
    notes.append(
        "Todas las celdas están en formato texto para que Sheets no altere GTIN ni popularity_rank. "
        "Si copias los datos a otra hoja, pon antes el formato en texto sin formato."
    )
    notes.extend(data.get("notas") or data.get("notes") or [])

    table = []
    for row in values:
        line = []
        for attr, kind, k in header_attrs:
            v = row[attr]
            if kind in REPEATED:
                v = v[k] if k < len(v) else ""
            line.append(v)
        table.append(line)
    return headers, table, notes


def write_xlsx(headers, table, notes, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Feed"
    bold = Font(bold=True)
    green = PatternFill("solid", fgColor="C6EFCE")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold
        cell.number_format = "@"
        if h in CONVERSATIONAL:
            cell.fill = green
    for r, line in enumerate(table, start=2):
        for c, v in enumerate(line, start=1):
            cell = ws.cell(row=r, column=c, value=v if v != "" else None)
            cell.number_format = "@"
    ws.freeze_panes = "B2"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 28

    wn = wb.create_sheet("Notas")
    wn.column_dimensions["A"].width = 140
    wn.cell(row=1, column=1, value="Notas del feed suplementario").font = bold
    for i, n in enumerate(notes, start=3):
        cell = wn.cell(row=i, column=1, value=n)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def write_sheets_json(headers, table, notes, out_path: Path):
    """Valores y peticiones batchUpdate para crear el feed con el conector de Google Sheets.

    Supone una hoja recién creada: la primera pestaña tiene sheetId 0.
    Las peticiones van ANTES de escribir valores: el formato texto tiene que existir
    antes de que entren GTIN y popularity_rank.
    """
    n_cols = len(headers)
    green = {"red": 0.776, "green": 0.937, "blue": 0.808}
    requests = [
        # Una hoja nueva trae 26 columnas y 1.000 filas: se amplía antes de formatear,
        # o cualquier petición sobre la columna 27+ falla y no se aplica nada del lote.
        {"updateSheetProperties": {
            "properties": {"sheetId": 0, "title": "Feed",
                           "gridProperties": {"frozenRowCount": 1,
                                              "columnCount": max(26, n_cols),
                                              "rowCount": max(1000, len(table) + 1)}},
            "fields": "title,gridProperties.frozenRowCount,"
                      "gridProperties.columnCount,gridProperties.rowCount"}},
        {"addSheet": {"properties": {"sheetId": 1, "title": "Notas"}}},
        {"repeatCell": {
            "range": {"sheetId": 0},
            "cell": {"userEnteredFormat": {"numberFormat": {"type": "TEXT"}}},
            "fields": "userEnteredFormat.numberFormat"}},
        {"repeatCell": {
            "range": {"sheetId": 0, "startRowIndex": 0, "endRowIndex": 1,
                      "startColumnIndex": 0, "endColumnIndex": n_cols},
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
            "fields": "userEnteredFormat.textFormat.bold"}},
        {"updateDimensionProperties": {
            "range": {"sheetId": 0, "dimension": "COLUMNS",
                      "startIndex": 0, "endIndex": n_cols},
            "properties": {"pixelSize": 200}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {
            "range": {"sheetId": 1, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 1},
            "properties": {"pixelSize": 900}, "fields": "pixelSize"}},
        {"repeatCell": {
            "range": {"sheetId": 1, "startColumnIndex": 0, "endColumnIndex": 1},
            "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
            "fields": "userEnteredFormat.wrapStrategy,userEnteredFormat.verticalAlignment"}},
    ]
    for c, h in enumerate(headers):
        if h in CONVERSATIONAL:
            requests.append({"repeatCell": {
                "range": {"sheetId": 0, "startRowIndex": 0, "endRowIndex": 1,
                          "startColumnIndex": c, "endColumnIndex": c + 1},
                "cell": {"userEnteredFormat": {"backgroundColor": green}},
                "fields": "userEnteredFormat.backgroundColor"}})

    notas = [["Notas del feed suplementario"], [""]] + [[n] for n in notes]
    payload = {
        "requests": requests,
        "feed_range": "Feed!A1",
        "feed_values": [headers] + table,
        "notas_range": "Notas!A1",
        "notas_values": notas,
        "filas": len(table),
        "columnas": n_cols,
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")


def main():
    args = sys.argv[1:]
    sheets_json = None
    if "--sheets-json" in args:
        i = args.index("--sheets-json")
        if i + 1 >= len(args):
            sys.exit(__doc__)
        sheets_json = Path(args[i + 1])
        del args[i:i + 2]
    if len(args) != 2:
        sys.exit(__doc__)
    src, dst = Path(args[0]), Path(args[1])
    if dst.suffix.lower() != ".xlsx":
        sys.exit("La salida tiene que ser un .xlsx")
    data = json.loads(src.read_text(encoding="utf-8"))
    try:
        headers, table, notes = build(data)
    except BuildError as e:
        sys.exit(f"ERROR: {e}")
    write_xlsx(headers, table, notes, dst)
    print(f"Feed generado: {dst}")
    if sheets_json:
        write_sheets_json(headers, table, notes, sheets_json)
        print(f"Datos para Google Sheets: {sheets_json}")
    print(f"Notas automáticas: {len(notes)}")


if __name__ == "__main__":
    main()
