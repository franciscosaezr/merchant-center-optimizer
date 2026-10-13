#!/usr/bin/env python3
"""
validate_feed.py — Valida un feed suplementario de Merchant Center en formato Google Sheets.

Uso:
    python validate_feed.py feed.xlsx            (lee la pestaña "Feed", o la primera)
    python validate_feed.py feed.csv | feed.tsv  (descarga de Google Sheets)

Comprueba límites de Google, formato de los atributos de grupo, escapado, GTIN,
structured_title/description, duplicados entre atributos y lenguaje promocional.

Salida: ERROR (hay que corregir) y AVISO (revisar). Código de salida 1 si hay errores.
"""
import csv
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

AI_RE = re.compile(r'^([a-z_]+):"(.*)"$', re.S)

REL_TYPES = {"part_of_set", "required_part", "often_bought_with",
             "substitute", "different_brand", "accessory"}
ID_TYPES = {"gtin", "id"}
PRIMARY_USUALLY_SENDS = {"title", "description", "brand", "gtin", "mpn",
                         "google_product_category", "product_type"}

PROMO = [
    r"gratis", r"env[ií]o", r"oferta", r"descuento", r"rebajas?", r"promoci[oó]n",
    r"compra (ya|ahora)", r"mejor precio", r"barat[oa]s?", r"liquidaci[oó]n",
    r"(el|la|los|las) mejor(es)?", r"n[uú]mero 1", r"free shipping", r"\bsale\b",
    r"discount", r"buy now", r"best price", r"cheap",
]
PROMO_RE = re.compile(r"\b(" + "|".join(PROMO) + r")\b", re.I)
EMOJI_RE = re.compile("[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F2FF]")
CAPS_RE = re.compile(r"\b[A-ZÁÉÍÓÚÑÜ]{5,}\b")
CAPS_OK = {"LED", "USB", "HDMI", "FULL", "WIFI", "NFC", "GPS", "VGA"}

STOP = set("""de la el los las un una unos unas y o en con sin para por que se su sus al del lo
como mas más muy es son esta este esto tiene tienen puede pueden ya no si sí cada hay
the a an and or of to in on with for is are it this that be can not yes""".split())


# ---------------------------------------------------------------- utilidades de parseo

def split_unescaped(text: str, sep: str) -> list:
    """Divide por `sep` ignorando separadores escapados con \\ o dentro de comillas dobles."""
    parts, buf, i, quoted = [], [], 0, False
    while i < len(text):
        ch = text[i]
        if ch == "\\" and i + 1 < len(text):
            buf.append(text[i:i + 2])
            i += 2
            continue
        if ch == '"':
            quoted = not quoted
        if ch == sep and not quoted:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
        i += 1
    parts.append("".join(buf))
    return parts


def unescape(text: str) -> str:
    out, i = [], 0
    while i < len(text):
        if text[i] == "\\" and i + 1 < len(text):
            out.append(text[i + 1])
            i += 2
        else:
            out.append(text[i])
            i += 1
    s = "".join(out).strip()
    if len(s) >= 2 and s[0] == s[-1] == '"':
        s = s[1:-1]
    return s


def tokens(text: str) -> set:
    t = unicodedata.normalize("NFKD", text.lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return {w for w in re.findall(r"[a-z0-9]+", t) if len(w) > 2 and w not in STOP}


def gtin_ok(g: str) -> bool:
    if not g.isdigit() or len(g) not in (8, 12, 13, 14):
        return False
    digits = [int(d) for d in g]
    check = digits.pop()
    total = sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(digits)))
    return (10 - total % 10) % 10 == check


# ---------------------------------------------------------------- lectura

def read_rows(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Falta openpyxl. Instálalo con: pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Feed"] if "Feed" in wb.sheetnames else wb.worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        delim = "\t" if suffix == ".tsv" else ","
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f, delimiter=delim))
    if not rows:
        sys.exit("Archivo vacío.")
    headers = [h.strip().lower() for h in rows[0]]
    products = []
    for r in rows[1:]:
        if not any(c.strip() for c in r):
            continue
        cols = defaultdict(list)
        for h, v in zip(headers, r):
            if h:
                cols[h].append(v.strip())
        products.append(cols)
    return headers, products


def values_of(cols, attr):
    """Todos los valores de un atributo repetido (varias columnas y/o varios por celda)."""
    out = []
    for cell in cols.get(attr, []):
        if cell:
            out.extend(p.strip() for p in split_unescaped(cell, ",") if p.strip())
    return out


def single(cols, attr):
    vals = [v for v in cols.get(attr, []) if v]
    return vals[0] if vals else ""


# ---------------------------------------------------------------- validación

class Report:
    def __init__(self):
        self.items = []

    def err(self, pid, msg):
        self.items.append(("ERROR", pid, msg))

    def warn(self, pid, msg):
        self.items.append(("AVISO", pid, msg))


def check_text(rep, pid, attr, text):
    if EMOJI_RE.search(text):
        rep.err(pid, f"{attr}: contiene emojis.")
    m = PROMO_RE.search(text)
    if m:
        rep.warn(pid, f"{attr}: posible lenguaje promocional ('{m.group(0)}').")
    if "%" in text:
        rep.warn(pid, f"{attr}: contiene '%'. Revisa que no sea una promesa o un descuento.")
    caps = [w for w in CAPS_RE.findall(text) if w not in CAPS_OK]
    if caps:
        rep.warn(pid, f"{attr}: palabras en mayúsculas ({', '.join(caps[:3])}). Solo valen siglas.")
    if "  " in text:
        rep.warn(pid, f"{attr}: espacios dobles.")


def check_ai(rep, pid, attr, raw, limit):
    m = AI_RE.match(raw)
    if not m:
        rep.err(pid, f"{attr}: formato inválido. Debe ser trained_algorithmic_media:\"texto\".")
        return ""
    source, content = m.group(1), m.group(2)
    if source != "trained_algorithmic_media":
        rep.warn(pid, f"{attr}: digital_source_type = '{source}'. Lo generado con IA va como trained_algorithmic_media.")
    if '"' in content:
        rep.err(pid, f"{attr}: comillas dobles dentro del texto.")
    if not 1 <= len(content) <= limit:
        rep.err(pid, f"{attr}: {len(content)} caracteres (límite {limit}).")
    check_text(rep, pid, attr, content)
    return content


def validate(headers, products):
    rep = Report()
    seen = set()
    groups = defaultdict(list)

    for idx, cols in enumerate(products, start=2):
        pid = single(cols, "id") or f"fila {idx}"

        # id
        if not single(cols, "id"):
            rep.err(pid, "Sin id.")
        elif pid in seen:
            rep.err(pid, "id duplicado.")
        elif len(pid) > 50:
            rep.err(pid, f"id de {len(pid)} caracteres (máx. 50).")
        seen.add(pid)

        # Celdas vacías en columnas que el principal suele enviar
        for attr in PRIMARY_USUALLY_SENDS:
            if attr in cols and not single(cols, attr):
                rep.warn(pid, f"{attr}: columna presente pero vacía. No está documentado si borra el valor del principal.")

        # Títulos
        title = single(cols, "title")
        st_raw = single(cols, "structured_title")
        st = check_ai(rep, pid, "structured_title", st_raw, 150) if st_raw else ""
        if title:
            if len(title) > 150:
                rep.err(pid, f"title: {len(title)} caracteres (máx. 150).")
            check_text(rep, pid, "title", title)
        if title and st_raw:
            rep.warn(pid, "Lleva title y structured_title: Google usará title.")
        main_title = st or title

        short = single(cols, "short_title")
        if short:
            if len(short) > 150:
                rep.err(pid, f"short_title: {len(short)} caracteres (máx. 150).")
            elif len(short) > 65:
                rep.warn(pid, f"short_title: {len(short)} caracteres (recomendado ≤ 65).")
            check_text(rep, pid, "short_title", short)

        # Descripciones
        desc = single(cols, "description")
        sd_raw = single(cols, "structured_description")
        sd = check_ai(rep, pid, "structured_description", sd_raw, 5000) if sd_raw else ""
        if desc:
            if len(desc) > 5000:
                rep.err(pid, f"description: {len(desc)} caracteres (máx. 5.000).")
            check_text(rep, pid, "description", desc)
        if desc and sd_raw:
            rep.warn(pid, "Lleva description y structured_description: Google usará description.")
        main_desc = sd or desc
        if main_title and main_desc and main_desc.lower().startswith(main_title.lower()[:60]):
            rep.warn(pid, "La descripción empieza repitiendo el título.")

        # Clasificación
        pt = single(cols, "product_type")
        if pt:
            if len(pt) > 750:
                rep.err(pid, f"product_type: {len(pt)} caracteres (máx. 750).")
            if "," in pt:
                rep.err(pid, "product_type: contiene comas (Google las toma como separador).")
            if ">" in pt and not re.search(r" > ", pt):
                rep.warn(pid, "product_type: separa los niveles con ' > ' (con espacios).")
        gpc = single(cols, "google_product_category")
        if gpc and not gpc.isdigit():
            rep.warn(pid, "google_product_category: mejor el ID numérico que la ruta.")

        # Identificadores
        gtin = single(cols, "gtin")
        if gtin and not gtin_ok(gtin):
            rep.err(pid, f"gtin '{gtin}' inválido (longitud o dígito de control).")

        # Highlights
        hl = [unescape(h) for h in values_of(cols, "product_highlight")]
        if hl:
            if len(hl) < 2:
                rep.err(pid, "product_highlight: mínimo 2.")
            elif len(hl) > 100:
                rep.err(pid, "product_highlight: máximo 100.")
            elif not 4 <= len(hl) <= 6:
                rep.warn(pid, f"product_highlight: {len(hl)} valores (recomendado 4–6).")
            for h in hl:
                if len(h) > 150:
                    rep.err(pid, f"product_highlight de {len(h)} caracteres (máx. 150): '{h[:50]}…'")
                check_text(rep, pid, "product_highlight", h)
            if len({x.lower() for x in hl}) < len(hl):
                rep.warn(pid, "product_highlight: valores repetidos.")

        # Product detail
        details = []
        for raw in values_of(cols, "product_detail"):
            parts = split_unescaped(raw, ":")
            if len(parts) != 3:
                rep.err(pid, f"product_detail con {len(parts) - 1} separadores ':' sin escapar (deben ser 2): '{raw[:60]}'")
                continue
            sec, name, val = (unescape(p) for p in parts)
            if not name or not val:
                rep.err(pid, f"product_detail sin nombre o valor: '{raw[:60]}'")
            if len(sec) > 140 or len(name) > 140:
                rep.err(pid, f"product_detail: sección o nombre > 140 caracteres: '{name[:40]}'")
            if len(val) > 1000:
                rep.err(pid, f"product_detail: valor > 1.000 caracteres: '{name}'")
            check_text(rep, pid, "product_detail", val)
            details.append(f"{name} {val}")

        # Q&A
        qas, total = [], 0
        for raw in values_of(cols, "question_and_answer"):
            parts = split_unescaped(raw, ":")
            if len(parts) != 2:
                rep.err(pid, f"question_and_answer con {len(parts) - 1} separadores ':' sin escapar (debe ser 1): '{raw[:60]}'")
                continue
            q, a = unescape(parts[0]), unescape(parts[1])
            if not q or not a:
                rep.err(pid, f"question_and_answer sin pregunta o respuesta: '{raw[:60]}'")
            if len(q) > 1000 or len(a) > 1000:
                rep.err(pid, f"question_and_answer: pregunta o respuesta > 1.000 caracteres: '{q[:40]}'")
            total += len(q) + len(a)
            check_text(rep, pid, "question_and_answer", a)
            qas.append((q, a))
        if len(qas) > 30:
            rep.err(pid, f"question_and_answer: {len(qas)} pares (máx. 30).")
        if total > 10000:
            rep.err(pid, f"question_and_answer: {total} caracteres en total (máx. 10.000).")
        qs = [tokens(q) for q, _ in qas]
        for i in range(len(qs)):
            for j in range(i + 1, len(qs)):
                if qs[i] and qs[j] and len(qs[i] & qs[j]) / len(qs[i] | qs[j]) >= 0.4:
                    rep.warn(pid, f"Preguntas parecidas: '{qas[i][0]}' / '{qas[j][0]}'. Revisa que no se solapen ni se contradigan.")

        # Duplicados Q&A vs resto de atributos
        pool = tokens(" ".join([main_title, main_desc, *hl, *details]))
        for q, a in qas:
            at = tokens(a)
            if len(at) >= 3 and len(at & pool) / len(at) >= 0.7:
                rep.warn(pid, f"La respuesta a '{q}' repite sobre todo datos que ya están en título, descripción, highlights o product_detail.")

        # Related product
        rel = values_of(cols, "related_product")
        if len(rel) > 30:
            rep.err(pid, f"related_product: {len(rel)} valores (máx. 30).")
        for raw in rel:
            parts = [unescape(p) for p in split_unescaped(raw, ":")]
            if len(parts) != 3:
                rep.err(pid, f"related_product debe ser tipo:tipo_id:identificador: '{raw}'")
                continue
            rt, it, ident = parts
            if rt not in REL_TYPES:
                rep.err(pid, f"related_product: relationship_type '{rt}' no válido.")
            if it not in ID_TYPES:
                rep.err(pid, f"related_product: identifier_type '{it}' no válido (gtin o id).")
            if not re.fullmatch(r"[A-Za-z0-9_-]+", ident):
                rep.err(pid, f"related_product: identificador '{ident}' con caracteres no permitidos.")
            if it == "gtin" and not gtin_ok(ident):
                rep.err(pid, f"related_product: GTIN '{ident}' inválido.")
            if ident == pid:
                rep.err(pid, "related_product apunta al propio producto.")

        # Variantes
        igt = single(cols, "item_group_title")
        vo = values_of(cols, "variant_option")
        names = set()
        for raw in vo:
            parts = split_unescaped(raw, ":")
            if len(parts) != 2:
                rep.err(pid, f"variant_option debe ser nombre:valor: '{raw}'")
                continue
            names.add(unescape(parts[0]).lower())
        if vo and not igt:
            rep.err(pid, "variant_option sin item_group_title.")
        if igt:
            groups[igt].append((pid, frozenset(names)))

        # Popularity rank
        pr = single(cols, "popularity_rank")
        if pr:
            if not re.fullmatch(r"\d{1,3}(\.\d)?", pr) or float(pr) > 100:
                rep.err(pid, f"popularity_rank '{pr}': número 0–100, máx. 1 decimal, punto decimal y sin %.")

        # Documentos y vídeos
        docs = values_of(cols, "document_link")
        if len(docs) > 5:
            rep.err(pid, f"document_link: {len(docs)} URLs (máx. 5).")
        for u in docs:
            if not re.match(r"https?://", u):
                rep.err(pid, f"document_link sin http(s): '{u}'")
            if len(u) > 2000 or not u.isascii() or " " in u:
                rep.err(pid, f"document_link: URL demasiado larga, con espacios o con caracteres no ASCII: '{u[:60]}'")
            if not u.lower().split("?")[0].endswith(".pdf"):
                rep.warn(pid, f"document_link no termina en .pdf; tiene que apuntar a un PDF: '{u[:60]}'")
        if docs and qas:
            rep.warn(pid, "Hay document_link y question_and_answer: Google pide no repetir en las FAQ lo que ya está en los PDF.")
        for u in values_of(cols, "video_link"):
            if not re.match(r"https?://", u):
                rep.err(pid, f"video_link sin http(s): '{u}'")
            if len(u) > 2000:
                rep.err(pid, "video_link > 2.000 caracteres.")
            if "youtube.com" in u or "youtu.be" in u:
                rep.warn(pid, "video_link de YouTube: la compatibilidad no está clara en la documentación.")

    for igt, members in groups.items():
        sets = {m[1] for m in members}
        if len(sets) > 1:
            rep.warn(", ".join(m[0] for m in members),
                     f"Grupo '{igt}': las variantes no usan el mismo conjunto de nombres en variant_option.")

    return rep


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    headers, products = read_rows(Path(sys.argv[1]))
    if "id" not in headers:
        sys.exit("ERROR: no hay columna 'id'.")
    rep = validate(headers, products)
    errors = [i for i in rep.items if i[0] == "ERROR"]
    warns = [i for i in rep.items if i[0] == "AVISO"]
    for level, pid, msg in errors + warns:
        print(f"{level:5} [{pid}] {msg}")
    print(f"\n{len(products)} productos · {len(errors)} errores · {len(warns)} avisos")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
